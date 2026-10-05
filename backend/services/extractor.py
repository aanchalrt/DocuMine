"""
Document text/table extraction service.
Supports PDF (pdfplumber + OCR fallback), DOCX, XLSX, and image files.
"""
from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------------------
# Optional heavy imports — handled gracefully so the module always loads
# ---------------------------------------------------------------------------

try:
    import pdfplumber  # type: ignore
except ImportError:
    pdfplumber = None  # type: ignore

try:
    from pdf2image import convert_from_path  # type: ignore
except ImportError:
    convert_from_path = None  # type: ignore

try:
    import pytesseract  # type: ignore
    from PIL import Image  # type: ignore
except ImportError:
    pytesseract = None  # type: ignore
    Image = None  # type: ignore

try:
    import docx  # python-docx  # type: ignore
except ImportError:
    docx = None  # type: ignore

try:
    import openpyxl  # type: ignore
except ImportError:
    openpyxl = None  # type: ignore


class DocumentExtractor:
    """Extracts text and tabular data from various document formats."""

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def extract(self, file_path: str, filename: str, mime_type: str) -> dict:
        """
        Extract text, tables, and document type from a file.

        Returns:
            {
                "text": str,
                "tables": list,
                "doc_type": str,
            }
        """
        ext = Path(filename).suffix.lower()
        text: str = ""
        tables: list = []

        try:
            if ext == ".docx":
                text, tables = self._extract_docx(file_path)
            elif ext == ".xlsx":
                text, tables = self._extract_xlsx(file_path)
            elif ext == ".pdf":
                text, tables = self._extract_pdf(file_path)
            elif ext in {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"}:
                text = self._extract_image_ocr(file_path)
                tables = []
            else:
                # Generic fallback — treat as plain text
                with open(file_path, "r", encoding="utf-8", errors="ignore") as fh:
                    text = fh.read()
        except Exception as exc:  # noqa: BLE001
            text = f"[Extraction error: {exc}]"

        doc_type = self._detect_doc_type(filename, text)
        return {"text": text, "tables": tables, "doc_type": doc_type}

    # ------------------------------------------------------------------
    # Private extraction helpers
    # ------------------------------------------------------------------

    def _extract_docx(self, file_path: str) -> tuple[str, list]:
        """Extract text and tables from a .docx file using python-docx."""
        if docx is None:
            raise RuntimeError("python-docx is not installed")

        document = docx.Document(file_path)
        paragraphs = [para.text for para in document.paragraphs if para.text.strip()]

        table_texts: list[list[list[str]]] = []
        for table in document.tables:
            tbl: list[list[str]] = []
            for row in table.rows:
                tbl.append([cell.text.strip() for cell in row.cells])
            table_texts.append(tbl)

        # Flatten table data into running text as well
        flat_table_text = "\n".join(
            " | ".join(cell for cell in row)
            for tbl in table_texts
            for row in tbl
            if any(cell for cell in row)
        )

        combined_text = "\n".join(paragraphs)
        if flat_table_text:
            combined_text += "\n\n" + flat_table_text

        return combined_text, table_texts

    def _extract_xlsx(self, file_path: str) -> tuple[str, list]:
        """Extract data from all sheets in an .xlsx workbook."""
        if openpyxl is None:
            raise RuntimeError("openpyxl is not installed")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        all_tables: list[dict[str, Any]] = []
        text_parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers: list[str] = []
            rows_as_dicts: list[dict] = []
            raw_rows: list[list[str]] = []

            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                str_row = [str(cell) if cell is not None else "" for cell in row]
                raw_rows.append(str_row)

                if row_idx == 0:
                    headers = str_row
                else:
                    row_dict = {
                        headers[i] if i < len(headers) else f"col_{i}": str_row[i]
                        for i in range(len(str_row))
                    }
                    rows_as_dicts.append(row_dict)

            all_tables.append(
                {
                    "sheet": sheet_name,
                    "headers": headers,
                    "rows": rows_as_dicts,
                    "raw": raw_rows,
                }
            )

            # Build a text representation for the sheet
            text_parts.append(f"=== Sheet: {sheet_name} ===")
            if headers:
                text_parts.append(" | ".join(headers))
            for row_dict in rows_as_dicts:
                text_parts.append(
                    " | ".join(f"{k}: {v}" for k, v in row_dict.items() if v)
                )

        return "\n".join(text_parts), all_tables

    def _extract_pdf(self, file_path: str) -> tuple[str, list]:
        """
        Extract text from a PDF.
        Falls back to OCR via pdf2image + pytesseract for scanned PDFs
        (heuristic: extracted text < 100 characters).
        """
        text = ""
        tables: list = []

        # --- Primary: pdfplumber ---
        if pdfplumber is not None:
            try:
                with pdfplumber.open(file_path) as pdf:
                    page_texts: list[str] = []
                    for page in pdf.pages:
                        page_text = page.extract_text() or ""
                        page_texts.append(page_text)

                        # pdfplumber table extraction
                        for tbl in page.extract_tables() or []:
                            tables.append(
                                [
                                    [str(cell) if cell is not None else "" for cell in row]
                                    for row in tbl
                                ]
                            )
                    text = "\n".join(page_texts)
            except Exception as exc:  # noqa: BLE001
                text = ""
                tables = []
                print(f"[extractor] pdfplumber error: {exc}")

        # --- Fallback: OCR ---
        if len(text.strip()) < 100:
            text = self._pdf_ocr_fallback(file_path)

        return text, tables

    def _pdf_ocr_fallback(self, file_path: str) -> str:
        """Convert PDF pages to images and run Tesseract OCR on each."""
        if convert_from_path is None or pytesseract is None:
            return "[OCR unavailable: pdf2image or pytesseract not installed]"

        try:
            with tempfile.TemporaryDirectory() as tmp_dir:
                images = convert_from_path(file_path, dpi=200, output_folder=tmp_dir)
                page_texts: list[str] = []
                for img in images:
                    page_text = pytesseract.image_to_string(img, lang="eng+hin")
                    page_texts.append(page_text)
                return "\n".join(page_texts)
        except Exception as exc:  # noqa: BLE001
            return f"[OCR error: {exc}]"

    def _extract_image_ocr(self, file_path: str) -> str:
        """Run Tesseract OCR directly on an image file."""
        if pytesseract is None or Image is None:
            return "[OCR unavailable: pytesseract or Pillow not installed]"

        try:
            img = Image.open(file_path)
            return pytesseract.image_to_string(img, lang="eng+hin")
        except Exception as exc:  # noqa: BLE001
            return f"[Image OCR error: {exc}]"

    # ------------------------------------------------------------------
    # Document type detection
    # ------------------------------------------------------------------

    def _detect_doc_type(self, filename: str, text: str) -> str:
        """
        Classify document type using filename and text heuristics.

        Priority order: filename keywords → text keywords → fallback.
        """
        name_lower = filename.lower()
        text_lower = text.lower()[:2000]  # Only examine start of text

        # Filename-first heuristics
        if "geological" in name_lower:
            return "geological_report"
        if "production" in name_lower:
            return "production_log"
        if "environmental" in name_lower or "compliance" in name_lower:
            return "compliance_report"
        if "correspondence" in name_lower or "noc" in name_lower or "clearance" in name_lower:
            return "correspondence"
        if "reserve" in name_lower or "archived" in name_lower or "2019" in name_lower:
            return "archived"

        # Text-based fallback heuristics
        if "geological" in text_lower or "stratigraphy" in text_lower or "lithology" in text_lower:
            return "geological_report"
        if "production log" in text_lower or "monthly production" in text_lower:
            return "production_log"
        if "compliance" in text_lower or "environmental clearance" in text_lower:
            return "compliance_report"
        if "no objection" in text_lower or "noc" in text_lower or "sender" in text_lower:
            return "correspondence"
        if "reserve" in text_lower or "archived" in text_lower:
            return "archived"

        return "general"
